import json
from pathlib import Path
from openpyxl import load_workbook
from docx import Document

base = Path(r"c:\Users\lenovo\FFA\insumos")
out = {}

# XLSX
xlsx = base / "Distribucion_HS_ROL_FFA_ECR_Salud.xlsx"
wb = load_workbook(xlsx, data_only=True)
sheets = {}
for sn in wb.sheetnames:
    ws = wb[sn]
    rows = []
    for row in ws.iter_rows(
        min_row=1, max_row=min(ws.max_row or 0, 200), values_only=True
    ):
        rows.append([str(c) if c is not None else "" for c in row])
    sheets[sn] = {"max_row": ws.max_row, "max_col": ws.max_column, "rows": rows}
out["xlsx"] = {"file": xlsx.name, "sheets": sheets}

# DOCX files
for docx_path in base.glob("*.docx"):
    doc = Document(docx_path)
    paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    tables = []
    for t in doc.tables:
        tbl = []
        for row in t.rows:
            tbl.append([cell.text.strip() for cell in row.cells])
        tables.append(tbl)
    out[docx_path.name] = {
        "paragraphs": paras,
        "tables": tables,
        "para_count": len(paras),
        "table_count": len(tables),
    }

digest_path = base / "_digest.json"
with open(digest_path, "w", encoding="utf-8") as f:
    json.dump(out, f, ensure_ascii=False, indent=2)

print("DIGEST SAVED:", digest_path)
print("XLSX sheets:", list(sheets.keys()))
for k, v in out.items():
    if k != "xlsx":
        print(f"DOCX: {k[:70]} | paras={v['para_count']} tables={v['table_count']}")
